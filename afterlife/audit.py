"""Hand audit of the papers that exactly one registry flags.

`afterlife audit sample` draws a seeded, stratified sample from the papers in
gold.registry that exactly one registry flags and writes audit/labels.csv for
a person to label. The sheet is blind: it never says which registry flagged a
paper, and a leading "RETRACTED:" is stripped from titles because it gives
both the source and the answer away. The key (sample id to paper and stratum)
goes only into the run report.

`afterlife audit score` reads the labels back, validates them against the
codebook in docs/audit-codebook.md, and writes per registry precision with
Wilson intervals plus a stratified estimate over the whole frame.
"""

from __future__ import annotations

import csv
import json
import math
import re
from pathlib import Path
from urllib.parse import quote

from afterlife.gold import ROOT

LABELS = ("retracted", "notice", "concern", "corrected", "not_retracted", "cannot_determine")
COLUMNS = [
    "sample_id",
    "title",
    "year",
    "doi_url",
    "pubmed_url",
    "pmc_url",
    "search_url",
    "label",
    "evidence_url",
    "notes",
]
SEED = "afterlife-audit-2026-10-26"
SHEET = ROOT / "audit" / "labels.csv"

_PREFIX = re.compile(
    r"^\s*(?:\[\s*(?:retracted|withdrawn)(?:\s+article)?\s*\]|(?:retracted|withdrawn)(?:\s+article)?\s*[:.\-])\s*",
    re.IGNORECASE,
)


def clean_title(title: str | None) -> str:
    return _PREFIX.sub("", title or "").strip()


def allocate(sizes: dict[str, int], n: int) -> dict[str, int]:
    """Equal allocation, capped at each stratum's size, the remainder spread
    over the strata that still have papers left."""
    alloc = dict.fromkeys(sizes, 0)
    open_ = sorted(s for s, size in sizes.items() if size > 0)
    left = n
    while left and open_:
        share = max(left // len(open_), 1)
        for s in list(open_):
            take = min(share, sizes[s] - alloc[s], left)
            alloc[s] += take
            left -= take
            if alloc[s] == sizes[s]:
                open_.remove(s)
            if not left:
                break
    return alloc


def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float] | None:
    if n == 0:
        return None
    p = k / n
    centre = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return (round(centre - half, 4), round(centre + half, 4))


def _frame(con) -> list[dict]:
    from afterlife.export import _lake
    from afterlife.findings import FLAGGED_SQL, _rows

    return _rows(
        con,
        f"""
        with e as (
          select *, {FLAGGED_SQL} as flagged
          from gold.registry where verdict in ('retracted', 'disputed')),
        oa as (
          select openalex_id, max(publication_year) as year
          from delta_scan('{_lake()}/silver/openalex_works') group by 1)
        select
          e.paper_key, e.doi, e.pmid, e.pmcid, e.title,
          coalesce(year(e.original_date), oa.year) as year,
          case when rw_status = 'retracted' then 'retraction_watch'
               when flag_openalex then 'openalex'
               when pmid is not null and flag_pubmed then 'pubmed'
               else 'pmc' end as stratum,
          md5('{SEED}' || e.paper_key) as draw
        from e left join oa using (openalex_id)
        where flagged = 1
        order by stratum, draw""",
    )


def _sheet_row(sample_id: str, r: dict) -> dict:
    title = clean_title(r["title"])
    return {
        "sample_id": sample_id,
        "title": title,
        "year": r["year"] or "",
        "doi_url": f"https://doi.org/{r['doi']}" if r["doi"] else "",
        "pubmed_url": f"https://pubmed.ncbi.nlm.nih.gov/{r['pmid']}/" if r["pmid"] else "",
        "pmc_url": f"https://pmc.ncbi.nlm.nih.gov/articles/{r['pmcid']}/" if r["pmcid"] else "",
        "search_url": "https://scholar.google.com/scholar?q=" + quote(f'"{title}"')
        if title
        else "",
        "label": "",
        "evidence_url": "",
        "notes": "",
    }


def sample(n: int = 200, sheet: Path = SHEET) -> dict:
    if sheet.exists():
        # Labels are hours of hand work; never overwrite them.
        raise SystemExit(f"{sheet} exists. Move it aside to draw a new sample.")
    from afterlife.export import _connect

    frame = _frame(_connect())
    sizes: dict[str, int] = {}
    for r in frame:
        sizes[r["stratum"]] = sizes.get(r["stratum"], 0) + 1
    alloc = allocate(sizes, n)
    taken = dict.fromkeys(sizes, 0)
    chosen = []
    for r in frame:  # ordered by stratum, then by the seeded hash
        if taken[r["stratum"]] < alloc[r["stratum"]]:
            taken[r["stratum"]] += 1
            chosen.append(r)
    # Shuffle across strata so the sheet's order says nothing about the source.
    chosen.sort(key=lambda r: r["draw"][::-1])
    # PubMed only papers have no title or year in gold (the landed PubMed set is
    # ids only), and a blank title would give the source away.
    from afterlife.sources import pubmed

    missing = [r["pmid"] for r in chosen if r["pmid"] and not (r["title"] and r["year"])]
    found = pubmed.summaries(missing)
    for r in chosen:
        if r["pmid"] in found:
            r["title"] = r["title"] or found[r["pmid"]]["title"]
            r["year"] = r["year"] or found[r["pmid"]]["year"]
    key = []
    sheet.parent.mkdir(parents=True, exist_ok=True)
    with sheet.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        for i, r in enumerate(chosen, 1):
            sid = f"A{i:03d}"
            w.writerow(_sheet_row(sid, r))
            key.append({"sample_id": sid, "paper_key": r["paper_key"], "stratum": r["stratum"]})
    return {
        "seed": SEED,
        "frame": "gold.registry, verdict retracted or disputed, flagged by exactly one registry",
        "frame_size": len(frame),
        "stratum_sizes": sizes,
        "allocation": alloc,
        "sheet": str(sheet.relative_to(ROOT)),
        "filled_from_pubmed_esummary": len(found),
        "titles_prefix_stripped": sum(
            clean_title(r["title"]) != (r["title"] or "").strip() for r in chosen
        ),
        "key": key,
    }


def score(labels: list[dict], sample_report: dict) -> dict:
    key = {k["sample_id"]: k["stratum"] for k in sample_report["key"]}
    sizes = sample_report["stratum_sizes"]
    bad = [
        f"{r['sample_id']}: {r['label']!r}"
        for r in labels
        if r["sample_id"] not in key
        or (r["label"].strip() and r["label"].strip().lower() not in LABELS)
    ]
    if bad:
        raise SystemExit("labels not in the codebook or unknown sample ids: " + ", ".join(bad))
    strata = {}
    for s in sizes:
        got = [
            r["label"].strip().lower()
            for r in labels
            if key[r["sample_id"]] == s and r["label"].strip()
        ]
        counts = {lab: got.count(lab) for lab in LABELS}
        decided = len(got) - counts["cannot_determine"]
        strata[s] = {
            "population": sizes[s],
            "sampled": sample_report["allocation"][s],
            "labelled": len(got),
            "counts": counts,
            "decided": decided,
            "precision": round(counts["retracted"] / decided, 4) if decided else None,
            "precision_ci95": wilson(counts["retracted"], decided),
            "notice_share": round(counts["notice"] / decided, 4) if decided else None,
        }
    complete = all(v["labelled"] == v["sampled"] for v in strata.values())
    est = None
    if complete and all(v["decided"] for v in strata.values() if v["sampled"]):
        # Stratified estimate with a finite population correction, so a stratum
        # labelled in full (a census) adds no sampling variance.
        total = sum(v["population"] for v in strata.values() if v["sampled"])
        p = var = 0.0
        for v in strata.values():
            if not v["sampled"]:
                continue
            w = v["population"] / total
            ph = v["precision"]
            p += w * ph
            fpc = 1 - v["decided"] / v["population"]
            var += w * w * ph * (1 - ph) / v["decided"] * fpc
        half = 1.96 * math.sqrt(var)
        est = {
            "precision": round(p, 4),
            "ci95": (round(max(p - half, 0), 4), round(min(p + half, 1), 4)),
            "frame_size": total,
        }
    return {"complete": complete, "strata": strata, "frame_estimate": est}


WORKBOOK = ROOT / "audit" / "labels.xlsx"
_HEADER = [
    "ID",
    "Title",
    "Year",
    "Open",
    "Label",
    "Evidence link",
    "Notes",
    "DOI",
    "PubMed",
    "PMC",
    "Search",
]
_FIRST = 6  # first data row; rows 1 to 4 are the legend, row 5 the header
_GUIDE = [
    ("How to label one row", None),
    ("1. Click Open. It goes to the publisher's page when there is one.", None),
    ("2. Decide from the publisher's page or the publisher's own notice.", None),
    ("3. Pick a label from the dropdown in the Label column.", None),
    ("4. Paste the page you used into Evidence link.", None),
    ("5. Add a note if something is odd. The row turns green when it has a label.", None),
    ("", None),
    ("The six labels", None),
    ("retracted", "The publisher retracted, withdrew, or removed it."),
    ("notice", "The item is itself the retraction notice, not the paper."),
    ("concern", "An expression of concern, and it is not retracted."),
    ("corrected", "A correction or erratum, and it is not retracted."),
    (
        "not_retracted",
        "You reached the publisher's page, it shows no retraction, and you found no notice.",
    ),
    ("cannot_determine", "You could not reach the publisher's page, or the evidence conflicts."),
    ("", None),
    ("Rules", None),
    (
        "Do not use Retraction Watch, OpenAlex, or the retraction banners on PubMed or PMC "
        "as evidence. They are what is being tested.",
        None,
    ),
    (
        "If Open goes to PubMed or PMC, use that page only to find the publisher's full text "
        "link, and decide from the publisher's page.",
        None,
    ),
    ("Do not open runs/*-audit-sample.json until every row is labelled.", None),
    ("About five minutes a row. After ten, choose cannot_determine and say what you tried.", None),
    ("", None),
    ("Notes tags (start Notes with any that apply)", None),
    (
        "front_matter  book_chapter  conference  retract_and_replace  partial  paywalled  "
        "non_english  dead_link  wrong_paper",
        None,
    ),
    ("", None),
    ("Example of a filled row (format only, not a real paper)", None),
    (
        "Label: notice    Evidence link: https://doi.org/10.xxxx/example    Notes: front_matter",
        None,
    ),
    ("", None),
    ("Full rules: docs/audit-codebook.md", None),
]


_GUIDE_HEADINGS = {
    "How to label one row",
    "The six labels",
    "Rules",
    "Notes tags (start Notes with any that apply)",
    "Example of a filled row (format only, not a real paper)",
}


def write_workbook(rows: list[dict], path: Path = WORKBOOK) -> dict:
    """The labelling sheet as an Excel workbook: one Open link per row, a
    dropdown limited to the six labels, and the rules on a Guide tab."""
    if path.exists():
        raise SystemExit(f"{path} exists. It may hold labels; move it aside first.")
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    link = Font(name="Arial", size=10, color="1F5FBF", underline="single")
    fill_in = PatternFill("solid", fgColor="FFF4C2")
    last = _FIRST + len(rows) - 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Label"
    ws["A1"] = "Hand audit: has the publisher really retracted each paper?"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = (
        "Fill in the yellow columns. Pick the label from the dropdown. "
        "The rules are on the Guide tab."
    )
    # A plain count beside a label: every spreadsheet app evaluates it the same way.
    ws["A3"] = f"=COUNTA(E{_FIRST}:E{last})"
    ws["B3"] = f"papers labelled so far, out of {len(rows)}"
    for c in ("A2", "A3", "B3"):
        ws[c].font = base
    ws["A3"].font = bold
    ws["A3"].alignment = Alignment(horizontal="left")
    for col, name in enumerate(_HEADER, 1):
        cell = ws.cell(row=_FIRST - 1, column=col, value=name)
        cell.font = bold
        if name in ("Label", "Evidence link", "Notes"):
            cell.fill = fill_in

    for i, r in enumerate(rows, _FIRST):
        best = next(
            (
                (u, t)
                for u, t in (
                    (r["doi_url"], "Open paper"),
                    (r["pubmed_url"], "Open in PubMed"),
                    (r["pmc_url"], "Open in PMC"),
                    (r["search_url"], "Search"),
                )
                if u
            ),
            ("", ""),
        )
        year = int(r["year"]) if str(r["year"]).isdigit() else r["year"]
        values = [
            r["sample_id"],
            r["title"],
            year,
            best[1],
            r.get("label", ""),
            r.get("evidence_url", ""),
            r.get("notes", ""),
            "",
            "",
            "",
            "",
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=i, column=col, value=v or None)
            cell.font = base
            cell.alignment = Alignment(vertical="top", wrap_text=col in (2, 7))
        for col, url in (
            (4, best[0]),
            (8, r["doi_url"]),
            (9, r["pubmed_url"]),
            (10, r["pmc_url"]),
            (11, r["search_url"]),
        ):
            if url:
                cell = ws.cell(row=i, column=col)
                cell.value = cell.value or _HEADER[col - 1]
                cell.hyperlink, cell.font = url, link
        for col in (5, 6, 7):
            ws.cell(row=i, column=col).fill = fill_in

    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(LABELS) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Not a label",
        error="Pick one of the six labels from the list.",
    )
    ws.add_data_validation(dv)
    dv.add(f"E{_FIRST}:E{last}")
    ws.conditional_formatting.add(
        f"A{_FIRST}:K{last}",
        FormulaRule(formula=[f"LEN($E{_FIRST})>0"], fill=PatternFill("solid", fgColor="DDF0DD")),
    )
    for col, width in zip("ABCDEFGHIJK", (7, 60, 6, 15, 18, 36, 28, 7, 9, 7, 9), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = f"B{_FIRST}"

    guide = wb.create_sheet("Guide")
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 80
    for i, (a, b) in enumerate(_GUIDE, 1):
        guide.cell(row=i, column=1, value=a or None)
        guide.cell(row=i, column=2, value=b)
        guide.cell(row=i, column=1).font = bold if a in _GUIDE_HEADINGS or b else base
        guide.cell(row=i, column=2).font = base
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    shown = path.relative_to(ROOT) if path.is_relative_to(ROOT) else path
    return {"workbook": str(shown), "rows": len(rows)}


def read_workbook(path: Path = WORKBOOK) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True)["Label"]
    out = []
    for row in ws.iter_rows(min_row=_FIRST, values_only=True):
        if row[0]:
            out.append(
                {
                    "sample_id": str(row[0]).strip(),
                    "label": str(row[4] or "").strip(),
                    "evidence_url": str(row[5] or "").strip(),
                    "notes": str(row[6] or "").strip(),
                }
            )
    return out


def make_sheet(sheet: Path = SHEET, path: Path = WORKBOOK) -> dict:
    with sheet.open(newline="") as fh:
        return write_workbook(list(csv.DictReader(fh)), path)


def run_score(
    sample_run: str | None = None, sheet: Path = SHEET, workbook: Path = WORKBOOK
) -> dict:
    path = (
        Path(sample_run) if sample_run else sorted((ROOT / "runs").glob("*-audit-sample.json"))[-1]
    )
    report = json.loads(path.read_text())["report"]
    with sheet.open(newline="") as fh:
        labels = list(csv.DictReader(fh))
    source = sheet
    if workbook.exists():
        # The workbook is where labelling happens; labels in both would be ambiguous.
        if any(r["label"].strip() for r in labels):
            raise SystemExit(f"labels found in both {sheet.name} and {workbook.name}; keep one")
        labels, source = read_workbook(workbook), workbook
    # The frame can shrink when gold is corrected (the 2026-10-29 notice fix moved
    # thousands of papers from one registry to two). Sampled papers still in the
    # frame, in the same stratum, remain a random sample of what is left there, so
    # score those against the current stratum sizes and keep the as drawn score.
    from afterlife.export import _connect

    current = {r["paper_key"]: r["stratum"] for r in _frame(_connect())}
    sizes: dict[str, int] = {}
    for s in current.values():
        sizes[s] = sizes.get(s, 0) + 1
    kept = [k for k in report["key"] if current.get(k["paper_key"]) == k["stratum"]]
    ids = {k["sample_id"] for k in kept}
    now = {
        "key": kept,
        "stratum_sizes": {s: sizes.get(s, 0) for s in report["stratum_sizes"]},
        "allocation": {s: sum(k["stratum"] == s for k in kept) for s in report["stratum_sizes"]},
    }
    return {
        "sample_run": path.name,
        "labels_from": source.name,
        "dropped_from_frame": len(report["key"]) - len(kept),
        **score([r for r in labels if r["sample_id"] in ids], now),
        "as_drawn": score(labels, report),
    }


SECOND = ROOT / "audit" / "labels_second.xlsx"


def second_sheet(n: int = 40, sheet: Path = SHEET, path: Path = SECOND) -> dict:
    """A blank workbook of n rows for a second labeller, drawn with a fixed seed
    from the committed blank sheet, so the first labels can never leak into it."""
    import hashlib

    with sheet.open(newline="") as fh:
        rows = list(csv.DictReader(fh))
    seed = SEED + "-second"
    pick = sorted(rows, key=lambda r: hashlib.md5((seed + r["sample_id"]).encode()).hexdigest())
    pick = sorted(pick[:n], key=lambda r: r["sample_id"])
    for r in pick:
        r.update(label="", evidence_url="", notes="")
    return {
        **write_workbook(pick, path),
        "seed": seed,
        "sample_ids": [r["sample_id"] for r in pick],
    }


def kappa(pairs: list[tuple[str, str]]) -> float | None:
    """Cohen's kappa for two raters; None when chance agreement is total."""
    if not pairs:
        return None
    n = len(pairs)
    po = sum(a == b for a, b in pairs) / n
    cats = {c for p in pairs for c in p}
    pe = sum(
        (sum(a == c for a, _ in pairs) / n) * (sum(b == c for _, b in pairs) / n) for c in cats
    )
    return None if pe == 1 else round((po - pe) / (1 - pe), 4)


def agreement(first: Path = WORKBOOK, second: Path = SECOND) -> dict:
    """Agreement between the first labeller and the second on the rows both labelled.
    The first labeller's labels stay the primary estimate; this only measures them."""
    a = {r["sample_id"]: r["label"].lower() for r in read_workbook(first)}
    b = {r["sample_id"]: r["label"].lower() for r in read_workbook(second) if r["label"]}
    bad = [f"{s}: {lab!r}" for s, lab in b.items() if lab not in LABELS or s not in a]
    if bad:
        raise SystemExit("second labels not in the codebook or unknown ids: " + ", ".join(bad))
    pairs = [(a[s], b[s]) for s in sorted(b)]
    binary = [(str(x == "retracted"), str(y == "retracted")) for x, y in pairs]
    return {
        "labelled_by_both": len(pairs),
        "percent_agreement": round(sum(x == y for x, y in pairs) / len(pairs), 4)
        if pairs
        else None,
        "kappa": kappa(pairs),
        "kappa_retracted_or_not": kappa(binary),
        "disagreements": [
            {"sample_id": s, "first": a[s], "second": b[s]} for s in sorted(b) if a[s] != b[s]
        ],
    }
